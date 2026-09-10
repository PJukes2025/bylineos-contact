#!/usr/bin/env python3
"""
Headless Edition-090 contact-sheet builder for GitHub Actions.
Reads the edition database + proof PDFs from Google Drive via a service account,
places each proof by the saved 'Pages Claimed' (table wins over filename),
keeps the latest proof per slug, rasterises with pdftoppm and renders the
self-contained HTML. Writes into bylineos-launcher-16up/contact/.

Env:
  GDRIVE_SA_KEY_PATH  path to the service-account JSON key
Outputs (relative to repo root):
  bylineos-launcher-16up/contact/ed090.html, index.html, current.html, .lastbuild
Prints DB_MODIFIED=<iso> and PAGES=<n>/68 for the workflow.
"""
import os, re, io, json, glob, shutil, subprocess, sys
from openpyxl import load_workbook
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
import gen

DB_ID   = "1LGXdD3U4gbowdM1A0wFHHOfpcrnaUVIRal4zoRiIfNQ"
PROOFS  = "1UvmjI1ArepHqsi-LFpqLKbwISJLspA70"
OUTDIR  = "bylineos-launcher-16up/contact"
ISSUE   = 68
WORK    = "work"

creds = service_account.Credentials.from_service_account_file(
    os.environ["GDRIVE_SA_KEY_PATH"],
    scopes=["https://www.googleapis.com/auth/drive.readonly"])
drive = build("drive", "v3", credentials=creds, cache_discovery=False)

def dl(file_id, path, export=None):
    req = (drive.files().export_media(fileId=file_id, mimeType=export) if export
           else drive.files().get_media(fileId=file_id))
    with open(path, "wb") as fh:
        d = MediaIoBaseDownload(fh, req, chunksize=8*1024*1024)
        done = False
        while not done:
            _, done = d.next_chunk()

if os.path.exists(WORK): shutil.rmtree(WORK)
os.makedirs(WORK+"/proofs"); os.makedirs(WORK+"/thumbs")

db_mod = drive.files().get(fileId=DB_ID, fields="modifiedTime").execute()["modifiedTime"]

# change-gate: skip the heavy build if the DB has not moved since last build
_marker=os.path.join(OUTDIR,".lastbuild")
if os.path.exists(_marker) and open(_marker).read().strip()==db_mod.strip():
    print("NOCHANGE db unchanged since last build:", db_mod); sys.exit(0)
dl(DB_ID, WORK+"/db.xlsx",
   export="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# ---- proofs list (exclude flatplans) ----
proofs=[]; tok=None
while True:
    r=drive.files().list(
        q=f"'{PROOFS}' in parents and mimeType='application/pdf' and trashed=false",
        fields="nextPageToken, files(id,name,modifiedTime)", pageSize=200, pageToken=tok).execute()
    proofs+=r.get("files",[]); tok=r.get("nextPageToken")
    if not tok: break
proofs=[f for f in proofs if "flatplan" not in f["name"].lower()]

# ---- DB parse ----
wb=load_workbook(WORK+"/db.xlsx", data_only=True)
ws=wb["Articles"] if "Articles" in wb.sheetnames else wb[wb.sheetnames[0]]
rows=list(ws.iter_rows(values_only=True))
hdr=[str(c).strip() if c is not None else "" for c in rows[0]]
def col(n):
    for i,h in enumerate(hdr):
        if h.lower()==n.lower(): return i
    return -1
C={k:col(k) for k in ["Headline","Section/Rubric","Status","Pages Claimed","Print Slug"]}
def pages(v):
    if v is None: return []
    if isinstance(v,(int,float)): return [int(v)]
    s=str(v).strip()
    m=re.match(r"^(\d+)(?:\.0)?\s*[–\-]\s*(\d+)(?:\.0)?$",s)
    if m: return list(range(int(m.group(1)),int(m.group(2))+1))
    m=re.match(r"^(\d+)(?:\.0)?$",s)
    return [int(m.group(1))] if m else []
arts=[]
for r in rows[1:]:
    if not any(r): continue
    pg=pages(r[C["Pages Claimed"]])
    if not pg: continue
    arts.append(dict(headline=str(r[C["Headline"]] or "").strip(),
        section=str(r[C["Section/Rubric"]] or "").strip(),
        status=str(r[C["Status"]] or "").strip(),
        pages=pg, slug=str(r[C["Print Slug"]] or "").strip().upper()))
slug2pages={a["slug"]:a["pages"] for a in arts if a["slug"]}

def vkey(t):
    s=t[:-4] if t.lower().endswith(".pdf") else t
    s=re.sub(r"^\s*\d+\s*[-–]\s*\d+\s*","",s)
    s=re.sub(r"^\s*\d+\s*","",s)
    s=re.sub(r"BT\.?90","",s,flags=re.I)
    s=re.sub(r"\((?:v)?\d+\)","",s,flags=re.I)
    s=re.sub(r"[_.\-–\s]+"," ",s).strip()
    return s.upper()
def fnpages(t):
    m=re.search(r"(\d+)\s*[-–]\s*(\d+)",t)
    if m: return list(range(int(m.group(1)),int(m.group(2))+1))
    m=re.match(r"^\s*(\d+)",t)
    return [int(m.group(1))] if m else []
def match(key,name):
    if key and key in slug2pages: return slug2pages[key]
    if key:
        for a in arts:
            if key and key in a["headline"].upper(): return a["pages"]
    return fnpages(name)

# latest per slug key
best={}
for f in proofs:
    k=vkey(f["name"]) or ("_fn_"+f["name"])
    if k not in best or f["modifiedTime"]>best[k]["modifiedTime"]: best[k]=f
chosen=list(best.values())

items=[]
for f in chosen:
    local=WORK+"/proofs/"+f["name"].replace("/","_")
    dl(f["id"], local)
    pg=match(vkey(f["name"]), f["name"])
    if pg: items.append((local,pg))
items.sort(key=lambda x:(len(x[1])==1, x[1][0]))
have=gen.rasterise_placed(items, WORK+"/thumbs", WORK+"/tmp")

meta={}
for a in arts:
    for p in a["pages"]: meta[int(p)]=dict(sec=a["section"],title=a["headline"],status=a["status"])

os.makedirs(OUTDIR, exist_ok=True)
subtitle="Rendered from the saved running order · page numbers from the page table · latest proof per slug · self-contained"
summ=[f"<b>{len(have)}</b>/{ISSUE} pages proofed", f"<b>{len(arts)}</b> articles placed"]
gen.render("Edition-090", OUTDIR+"/ed090.html", WORK+"/thumbs", meta,
           "Byline Times — Edition 090 · contact sheet", subtitle, summ,
           issue_pages=ISSUE, status_legend=True)

# current.html -> ed090
open(OUTDIR+"/current.html","w",encoding="utf-8").write(
 '<!DOCTYPE html><html lang="en"><head><meta charset="utf-8">'
 '<meta name="robots" content="noindex,nofollow,noarchive">'
 '<title>Current contact sheet — redirecting</title>'
 '<meta http-equiv="refresh" content="0; url=ed090.html">'
 '<link rel="canonical" href="ed090.html"></head>'
 '<body><p>Opening the latest contact sheet (Edition 90)… '
 '<a href="ed090.html">continue</a>.</p></body></html>')

# index.html count refresh (only touch the ed090 line's page count)
idx=OUTDIR+"/index.html"
if os.path.exists(idx):
    html=open(idx,encoding="utf-8").read()
    html=re.sub(r"(Edition 90[\s\S]{0,400}?)(\d+) / 68 pages proofed",
                lambda m: m.group(1)+f"{len(have)} / 68 pages proofed", html, count=1)
    open(idx,"w",encoding="utf-8").write(html)

open(OUTDIR+"/.lastbuild","w").write(db_mod)
print("DB_MODIFIED="+db_mod)
print(f"PAGES={len(have)}/{ISSUE}")
